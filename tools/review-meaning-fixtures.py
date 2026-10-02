#!/usr/bin/env python3
"""Build a synthetic One Page review docx from the approved template (issue 1242).

  python3 tools/review-meaning-fixtures.py SPEC.json OUT.docx

SPEC.json: {"direct", "start", "end", "core", "strengths": [...], "weaknesses": [...], "guidance": [...]}.
The text is invented: the Manager Tools Bob examples and made-up defects. No real review text
belongs in this public repo; tools/review-meaning.test.js is the only caller.
"""
import json
import os
import sys

from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "aac-skills", "aac-performance-review-audit",
                        "DIRECT NAME - Annual Performance Review - YEAR (template, 2026-09-23).docx")


def main(spec_path, out):
    spec = json.load(open(spec_path, encoding="utf-8"))
    d = Document(TEMPLATE)
    paras = d.paragraphs
    name = spec["direct"]
    paras[0].text = f"{name}’s Annual Performance Review"
    paras[1].text = "Delivered by: Pat Example  Date: <DATE OF REVIEW>"
    paras[2].text = "Method of delivery: In person"
    paras[3].text = (f"Last Review Date: {spec['start']}  Dates covered by this review: "
                     f"{spec['start']} – {spec['end']}")
    paras[4].text = "Core Message: " + spec["core"]
    slots = [p for p in paras if p.style.name == "List Paragraph"]
    for p, g in zip(slots, spec["guidance"]):
        p.text = g
    for p in slots[len(spec["guidance"]):]:
        p._element.getparent().remove(p._element)
    table = next(t for t in d.tables if "strength" in t.rows[0].cells[0].text.lower())
    for ci, items in ((0, spec["strengths"]), (1, spec["weaknesses"])):
        for r, text in enumerate(items, 1):
            table.rows[r].cells[ci].paragraphs[0].text = text
    d.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
