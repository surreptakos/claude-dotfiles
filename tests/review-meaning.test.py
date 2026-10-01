#!/usr/bin/env python3
"""The review scripts' `meaning` command at the CLI (issue 1243). Run:  python3 tests/review-meaning.test.py

Needs python-docx. Every fixture is synthetic: the approved review template filled with the
Manager Tools Bob examples and invented defects, written to a temp folder. No real review text.
TYPESAFE_JEV_STUB=off keeps Jev off the network. The copies of review_meaning.py and jev.py in the
two skill folders are pinned identical by tools/review-meaning-copies.test.js.
"""

import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
AUDIT = ROOT / "aac-skills" / "aac-performance-review-audit"
SELF = ROOT / "aac-skills" / "aac-review-self-check"
HOSTS = {
    "audit": [str(AUDIT / "review_gate_tools.py"), "meaning"],
    "self-check": [str(SELF / "review_format_check.py"), "meaning"],
}
TEMPLATE = next(AUDIT.glob("DIRECT NAME - Annual Performance Review - YEAR (template, *).docx"))
ENV = dict(os.environ, TYPESAFE_JEV_STUB="off", PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1")

BOB_SEER = ("Bob is my best customer service rep. He consistently exceeds every standard. "
            "He recently saved a difficult call after three other reps had failed. "
            "He's an example we ought to put on training videos.")
BOB_SUMEX = ("Bob is my best customer service rep. "
             "Recently he saved a difficult call despite 3 other reps not being able to.")
WEAKNESS = "Bob logs his call notes late. On 3/2/2026 he entered a full week of notes in one sitting."
GUIDANCE = ["Log every call note the same day. Help the new reps do the same.",
            "Lead the Friday call review with me."]


def build(path, strengths=(BOB_SEER, BOB_SUMEX), weaknesses=(WEAKNESS,), guidance=GUIDANCE):
    d = Document(str(TEMPLATE))
    paras = d.paragraphs
    paras[0].text = "Bob's Annual Performance Review"
    paras[3].text = "Last Review Date: 7/23/2025  Dates covered by this review: 7/24/2025 - 7/23/2026"
    paras[4].text = ("Core Message: Bob's results have met expectations since his last review. I am recommending "
                     "he maintain his current role and responsibilities at this time.")
    bullets = [p for p in paras if p.text.startswith("Start with a verb")]
    for p, text in zip(bullets, list(guidance) + [""] * len(bullets)):
        p.text = text
    table = d.tables[0]
    for col, items in ((0, strengths), (1, weaknesses)):
        for row, text in zip(table.rows[1:], items):
            row.cells[col].paragraphs[0].text = text
    d.save(str(path))
    return path


class Meaning(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.n = 0

    def review(self, **kw):
        self.n += 1
        return build(Path(self.tmp.name) / f"review-{self.n}.docx", **kw)

    def run_meaning(self, path, host="audit", direct="Bob"):
        cmd = [sys.executable, *HOSTS[host], str(path), "--end", "7/23/2026", "--start", "7/24/2025",
               "--direct", direct]
        r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", env=ENV, timeout=120)
        return r.returncode, r.stdout

    def fixes(self, out):
        return [line[4:] for line in out.splitlines() if line.startswith("  - ")]

    def test_clean_review_passes_from_both_folders(self):
        path = self.review()
        outs = []
        for host in HOSTS:
            code, out = self.run_meaning(path, host)
            self.assertEqual(code, 0, out)
            self.assertIn("MEANING: PASS", out)
            self.assertEqual(self.fixes(out), [])
            outs.append(out)
        self.assertEqual(outs[0], outs[1])

    def test_guidance_sentence_not_an_instruction_fails_from_both_folders(self):
        cases = {
            "He will log every call note the same day.": "He",
            "She logs call notes late.": "She",
            "They are behind on notes.": "They",
            "Bob's notes run a week late.": "Bob's",
            "I will review his notes each Friday.": "I will",
            "I'll meet with him every month.": "I'll",
            "I want to see same-day notes.": "I want",
        }
        for sentence, opener in cases.items():
            with self.subTest(sentence=sentence):
                path = self.review(guidance=["Log every call note the same day. " + sentence, GUIDANCE[1]])
                for host in HOSTS:
                    code, out = self.run_meaning(path, host)
                    self.assertEqual(code, 1, out)
                    self.assertIn("MEANING: FAIL (1 fix)", out)
                    self.assertEqual(self.fixes(out), [
                        f'Guidance 1: sentence 2 opens on "{opener}" and is not an instruction; '
                        "it must be an instruction for next year."])

    def test_seer_sentence_3_with_two_dates_fails_one_example(self):
        one_date = BOB_SEER.replace("He recently saved a difficult call",
                                    "On 3/2/2026 he saved a difficult call")
        code, out = self.run_meaning(self.review(strengths=(one_date,)))
        self.assertEqual(code, 0, out)

        two_dates = BOB_SEER.replace("He recently saved a difficult call after three other reps had failed.",
                                     "He saved a difficult call on 3/2/2026 and another on March 9, 2026.")
        path = self.review(strengths=(BOB_SUMEX, two_dates))
        for host in HOSTS:
            code, out = self.run_meaning(path, host)
            self.assertEqual(code, 1, out)
            self.assertEqual(self.fixes(out), [
                "S2: sentence 3 names 2 dates (3/2/2026, March 9, 2026); it must be one example."])

    def test_comma_counts_are_notes_and_never_change_the_exit_code(self):
        commas = BOB_SUMEX.replace("Recently he saved", "Recently, on a busy Monday, he saved")
        code, out = self.run_meaning(self.review(strengths=(BOB_SEER, commas)))
        self.assertEqual(code, 0, out)
        self.assertIn("Note: commas per item, keep them to a minimum: S1 0, S2 2, W1 0", out)
        self.assertEqual(self.fixes(out), [])

        code, out = self.run_meaning(self.review(strengths=(BOB_SEER, commas),
                                                 guidance=["He logs notes late.", GUIDANCE[1]]))
        self.assertEqual(code, 1, out)
        self.assertEqual(len(self.fixes(out)), 1)
        self.assertIn("S2 2", out)

    def test_could_not_check_exits_2_from_both_folders(self):
        bad = Path(self.tmp.name) / "not-a-review.docx"
        bad.write_text("not a docx", encoding="utf-8")
        for host in HOSTS:
            code, out = self.run_meaning(bad, host)
            self.assertEqual(code, 2, out)
            self.assertIn("MEANING: COULD NOT CHECK", out)
            self.assertIn("This is not a pass.", out)


if __name__ == "__main__":
    unittest.main()
