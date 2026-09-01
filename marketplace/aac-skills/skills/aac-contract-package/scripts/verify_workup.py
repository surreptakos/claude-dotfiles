"""Workup arithmetic checks — the mechanical half of PROMPT.md's workup rules.

Usage:  python verify_workup.py "<job folder or workup xlsx>" [--quiet]
Exit:   0 = no FAILs, 1 = at least one FAIL, 2 = bad input, 3 = no workup found.

Reads only. Never writes to the job folder and never saves a workbook.

What it checks, from PROMPT.md ("What you fix versus what you ask" and
Recurring Defect Checks, Workup arithmetic):

  1. Overwritten extension formulas — a typed number sitting in a column that
     otherwise holds row formulas (where `=C*D` belongs).
  2. Typed cost on a blank-quantity line — cost surviving where the intact
     formula would have yielded zero.
  3. SUM totals recomputed from their ranges — a total that does not follow
     from its parts, and costs sitting just outside a SUM's range.
  4. Stranded costs — a computed or typed amount no formula anywhere on the
     tab picks up, which is how a lower section's cost misses the tab total.

And the rule the checks must respect: a blank quantity with its formula intact
is an unused template line, not a defect. It is only reported when a lower
section prices that line (case 4 catches the stranded cost, not the blank).

Every finding carries the tab, the cell, and the amount, because that is what
the workup block of the review email needs. Findings are defect reports for
the rep and the estimator — never corrected prices.

Verifier discipline applies (SKILL.md): when a finding looks surprising, check
it against the workbook before reporting it. A tool bug reads exactly like a
package defect. Hidden and very-hidden tabs are skipped, per the File Reading
Rules; skipped tab names are listed so the skip is visible.
"""
import sys, os, re, glob, warnings
from collections import defaultdict
warnings.filterwarnings('ignore')

MIN_FORMULA_SHARE = 0.6   # a column is an "extension column" when >=60% of its
                          # populated numeric-ish cells are formulas
MIN_COLUMN_FORMULAS = 3   # ...and it holds at least this many formulas
MONEY_FLOOR = 0.005       # ignore zero/rounding dust when hunting strays

results = []
def rec(st, item, detail=''):
    results.append((st, item, detail))

CELL = re.compile(r'(\$?)([A-Z]{1,3})(\$?)(\d+)')
RANGE = re.compile(r'\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)')

def col_to_n(c):
    n = 0
    for ch in c:
        n = n * 26 + ord(ch) - 64
    return n

def n_to_col(n):
    s = ''
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s

def refs_of(formula):
    """Every cell a formula touches, ranges expanded. Same-sheet refs only;
    a cross-sheet ref is kept as ('SHEET', coord) so we can see tab links."""
    out = set()
    f = formula
    # cross-sheet refs: 'Tab Name'!A1 or Tab!A1:B9 — record and blank them out
    for m in re.finditer(r"(?:'([^']+)'|(\w+))!(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)", f):
        sheet = m.group(1) or m.group(2)
        for c, r in expand(m.group(3)):
            out.add((sheet, f'{c}{r}'))
    f = re.sub(r"(?:'[^']+'|\w+)!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?", ' ', f)
    for m in RANGE.finditer(f):
        c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        for cn in range(col_to_n(c1), col_to_n(c2) + 1):
            for r in range(r1, min(r2, r1 + 5000) + 1):
                out.add((None, f'{n_to_col(cn)}{r}'))
    f = RANGE.sub(' ', f)
    for m in CELL.finditer(f):
        out.add((None, f'{m.group(2)}{m.group(4)}'))
    return out

def expand(rng):
    m = RANGE.fullmatch(rng.replace('$', ''))
    if not m:
        m2 = re.fullmatch(r'([A-Z]{1,3})(\d+)', rng.replace('$', ''))
        return [(m2.group(1), int(m2.group(2)))] if m2 else []
    c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
    return [(n_to_col(cn), r) for cn in range(col_to_n(c1), col_to_n(c2) + 1)
            for r in range(r1, r2 + 1)]

def money(v):
    return float(v) if isinstance(v, (int, float)) else None


def check_tab(ws, wsv):
    tab = ws.title
    formulas = {}     # coord -> formula string
    numbers = {}      # coord -> literal numeric value
    values = {}       # coord -> cached/computed value
    by_col = defaultdict(lambda: {'f': [], 'n': []})
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if v is None:
                continue
            coord = c.coordinate
            if isinstance(v, str) and v.startswith('='):
                formulas[coord] = v
                by_col[c.column_letter]['f'].append(c.row)
                values[coord] = money(wsv[coord].value)
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                numbers[coord] = float(v)
                by_col[c.column_letter]['n'].append(c.row)
                values[coord] = float(v)

    # which cells does any formula on this tab consume?
    consumed = set()
    for coord, f in formulas.items():
        for sheet, ref in refs_of(f):
            if sheet is None or sheet == tab:
                consumed.add(ref)
    # cells consumed by other tabs' formulas are added by the caller

    # ---- 1. overwritten extension formulas ----
    ext_cols = {}
    for col, d in by_col.items():
        nf, nn = len(d['f']), len(d['n'])
        if nf >= MIN_COLUMN_FORMULAS and nf / (nf + nn) >= MIN_FORMULA_SHARE:
            ext_cols[col] = d
    reported = set()
    for col, d in ext_cols.items():
        lo, hi = min(d['f']), max(d['f'])
        strays = [r for r in d['n'] if lo <= r <= hi]
        # learn the factor columns from this column's intact formulas (=C5*D5)
        factor_cols = set()
        for r in d['f']:
            f = formulas[f'{col}{r}'].replace('$', '')
            m = re.fullmatch(r'=([A-Z]{1,3})(\d+)\*([A-Z]{1,3})(\d+)', f.replace(' ', ''))
            if m and int(m.group(2)) == r and int(m.group(4)) == r:
                factor_cols.update([m.group(1), m.group(3)])
        for r in strays:
            amt = numbers[f'{col}{r}']
            blanks = [fc for fc in sorted(factor_cols)
                      if ws[f'{fc}{r}'].value in (None, '')]
            # ---- 2. typed cost on a blank-quantity line (the stronger call) ----
            if blanks:
                rec('FAIL', f'{tab}: typed cost on a blank-quantity line',
                    f'{col}{r} holds {amt:,.2f} while {", ".join(f"{b}{r}" for b in blanks)} '
                    f'is blank — an intact formula would have yielded zero')
            else:
                rec('FAIL', f'{tab}: typed number where a formula belongs',
                    f'{col}{r} holds {amt:,.2f}; rows {lo}-{hi} of column {col} carry formulas')
            reported.add(f'{col}{r}')

    # ---- 3. SUM totals recomputed, and costs just outside the range ----
    for coord, f in formulas.items():
        m = re.fullmatch(r'=SUM\(\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)\)',
                         f.replace(' ', ''), re.I)
        if not m or m.group(1) != m.group(3):
            continue
        col, r1, r2 = m.group(1), int(m.group(2)), int(m.group(4))
        total = sum(values.get(f'{col}{r}') or 0 for r in range(r1, r2 + 1))
        cached = money(wsv[coord].value)
        if cached is not None and abs(cached - total) > 0.01:
            rec('WARN', f'{tab}: stale cached total',
                f'{coord} caches {cached:,.2f}; its range recomputes to {total:,.2f} '
                f'(Excel refreshes this on open — trust the recomputed figure)')
        # a value in the same column, within 3 rows above the range or between
        # the range end and the total, that the SUM does not reach
        tr = int(CELL.search(coord).group(4))
        for r in list(range(max(1, r1 - 3), r1)) + list(range(r2 + 1, tr)):
            cc = f'{col}{r}'
            amt = values.get(cc)
            if amt and abs(amt) > MONEY_FLOOR and cc not in consumed:
                rec('FAIL', f'{tab}: cost outside the total’s range',
                    f'{cc} holds {amt:,.2f}; {coord} sums {col}{r1}:{col}{r2} and misses it')
                reported.add(cc)

    return formulas, numbers, values, consumed, ext_cols, reported


def verify(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)
    visible = [ws for ws in wb.worksheets if ws.sheet_state == 'visible']
    skipped = [ws.title for ws in wb.worksheets if ws.sheet_state != 'visible']
    if skipped:
        rec('SKIP', 'Hidden tabs not read (File Reading Rules)', ', '.join(skipped))
    per_tab = {}
    for ws in visible:
        per_tab[ws.title] = check_tab(ws, wbv[ws.title])

    # cross-tab consumption, then the stranded-cost pass
    consumed_by_tab = defaultdict(set)
    for tab, (formulas, _, _, consumed, _, _) in per_tab.items():
        consumed_by_tab[tab] |= consumed
        for coord, f in formulas.items():
            for sheet, ref in refs_of(f):
                if sheet and sheet != tab:
                    consumed_by_tab[sheet].add(ref)

    for tab, (formulas, numbers, values, _, ext_cols, reported) in per_tab.items():
        consumed = consumed_by_tab[tab]
        for col, d in ext_cols.items():
            for r in sorted(d['f'] + d['n']):
                c = f'{col}{r}'
                if c in consumed or c in reported or (values.get(c) or 0) <= MONEY_FLOOR:
                    continue
                # a SUM formula nothing consumes is a total — totals are
                # allowed to be leaves. A literal or row formula is not.
                f = formulas.get(c, '')
                if re.match(r'=\s*SUM\(', f, re.I):
                    continue
                rec('WARN', f'{tab}: amount no formula picks up',
                    f'{c} holds {values[c]:,.2f} and nothing on any visible tab sums it '
                    f'— if a lower section priced this, its cost is stranded')


def find_workup(job):
    hits = []
    for pat in ('*Work*Up*.xls*', '*Workup*.xls*', '*work up*.xls*', '*WU*.xls*'):
        for base in (job, os.path.join(job, '*')):
            hits += glob.glob(os.path.join(base, pat))
    hits = [h for h in sorted(set(hits), key=os.path.getmtime, reverse=True)
            if not os.path.basename(h).startswith('~$')
            and os.path.splitext(h)[1].lower() in ('.xlsx', '.xlsm', '.xltx')
            and 'template' not in os.path.basename(h).lower()
            and 'equip & s' not in os.path.basename(h).lower()
            and not any(p in ('_extract', '_to_delete') for p in h.split(os.sep))]
    return hits


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    quiet = '--quiet' in sys.argv
    if not args:
        print(__doc__); sys.exit(2)
    target = args[0].rstrip('\\/')
    if os.path.isdir(target):
        wus = find_workup(target)
        if not wus:
            print('no workup workbook found in', target); sys.exit(3)
        if len(wus) > 1:
            rec('WARN', 'Multiple workups in the folder',
                f'newest used: {os.path.basename(wus[0])} — reconcile revisions line by '
                f'line per SCHEDULE-GENERATION-PROCEDURE section 1 before trusting either')
        target = wus[0]
    if not os.path.isfile(target):
        print('not a file:', target); sys.exit(2)
    print('=' * 78)
    print('WORKUP ARITHMETIC —', os.path.basename(target))
    print('=' * 78)
    try:
        verify(target)
    except Exception as e:
        rec('FAIL', 'Workup verifier ran to completion', f'{type(e).__name__}: {e}')
    for st in ('FAIL', 'WARN', 'SKIP', 'PASS'):
        if quiet and st in ('PASS', 'SKIP'):
            continue
        rows = [r for r in results if r[0] == st]
        if not rows:
            continue
        print(f'\n{st}  ({len(rows)})')
        for _, item, detail in rows:
            print(f'   {item}' + (f'  —  {detail}' if detail else ''))
    nf = sum(1 for r in results if r[0] == 'FAIL')
    nw = sum(1 for r in results if r[0] == 'WARN')
    print('\n' + '-' * 78)
    print(f'{nf} fail, {nw} warn.  Findings are defect reports carrying tab, cell and')
    print('amount — never corrected prices. Check surprising findings against the')
    print('workbook before reporting them; a tool bug reads exactly like a defect.')
    print('Labor hours, rates, markup and adequacy are never findings (PROMPT.md).')
    sys.exit(1 if nf else 0)


if __name__ == '__main__':
    main()
