"""Zoho CRM cross-checks for verify_package.py (issue 229; spec 215 stream D,
wave 2). Advisory only.

Zoho is deal-envelope metadata that gets validated against the package, never
trusted (CLAUDE.md hard rule 5). Every check here emits WARN or PASS with the
package value and the Zoho value both in the detail, or SKIP with a reason.
Nothing here ever FAILs, and nothing here writes: no Zoho value is copied into
the job folder, the schedule, the agreements or ``_facts.json``.

Credential: ZOHO_CLIENT_ID, ZOHO_CLIENT_SECRET and ZOHO_REFRESH_TOKEN, plus the
optional ZOHO_API_DOMAIN (default https://www.zohoapis.com) and
ZOHO_ACCOUNTS_DOMAIN (default derived from the API domain). Any of the three
required variables unset, or the network unreachable, degrades every check to
SKIP with the reason. The only live calls are an OAuth token refresh and one
read-only Deals search by Prospect #.

AAC_ZOHO_REPLAY names a recorded Deals-search response (JSON) that replaces
the live call; the test suite uses it with synthetic values so no test ever
reaches the live API. The credential gate still applies in replay mode.
"""
import json
import os
import re

CRED_VARS = ('ZOHO_CLIENT_ID', 'ZOHO_CLIENT_SECRET', 'ZOHO_REFRESH_TOKEN')
REPLAY_VAR = 'AAC_ZOHO_REPLAY'
TIMEOUT = 15

C_CONTACT = 'Zoho contact phone and email match the record'
C_DEAL = 'Zoho deal name, markup and subcontractor flag match the package'
C_PROSPECT = 'Zoho deal exists for the Prospect #'
C_MASTER = 'Zoho master-sent checkboxes match the Package situation'
C_TERM = 'Zoho effective agreement term set on a project carrying RMR'
CHECKS = (C_PROSPECT, C_CONTACT, C_DEAL, C_MASTER, C_TERM)

# Zoho Deals checkbox per Contract family (families are the SOW-BASELINES.md
# §3 column, read at run time through build_package.system_families()).
MASTER_BOXES = {
    'Commercial Fire': 'Sending_New_Fire_Master',
    'Commercial Security': 'Sending_New_Security_Master',
    'Elevator Monitoring': 'Sending_New_Elevator_Monitoring_Master_Agmt',
}
MARKUP_TOLERANCE_PTS = 0.5


def missing_credentials():
    return [v for v in CRED_VARS if not os.environ.get(v)]


def _api_domain():
    return (os.environ.get('ZOHO_API_DOMAIN') or 'https://www.zohoapis.com').rstrip('/')


def _accounts_domain():
    acc = os.environ.get('ZOHO_ACCOUNTS_DOMAIN')
    if acc:
        return acc.rstrip('/')
    m = re.search(r'zohoapis\.([a-z.]+)$', _api_domain())
    return f'https://accounts.zoho.{m.group(1) if m else "com"}'


def _access_token():
    import urllib.parse
    import urllib.request
    q = urllib.parse.urlencode({
        'refresh_token': os.environ['ZOHO_REFRESH_TOKEN'],
        'client_id': os.environ['ZOHO_CLIENT_ID'],
        'client_secret': os.environ['ZOHO_CLIENT_SECRET'],
        'grant_type': 'refresh_token',
    }).encode()
    req = urllib.request.Request(f'{_accounts_domain()}/oauth/v2/token', data=q)
    with urllib.request.urlopen(req, timeout=TIMEOUT) as r:
        tok = json.loads(r.read() or b'{}').get('access_token')
    if not tok:
        raise ValueError('token refresh returned no access_token')
    return tok


def _search(prospect):
    """Return (parsed_response, source_label). Raises on network trouble."""
    replay = os.environ.get(REPLAY_VAR)
    if replay:
        with open(replay, encoding='utf-8') as f:
            return json.load(f), f'recorded response {os.path.basename(replay)}'
    import urllib.parse
    import urllib.request
    token = _access_token()
    url = (f'{_api_domain()}/crm/v8/Deals/search?' + urllib.parse.urlencode(
        {'criteria': f'(True_Lead_Number:equals:{prospect})'}))
    req = urllib.request.Request(
        url, headers={'Authorization': f'Zoho-oauthtoken {token}'})
    with urllib.request.urlopen(req, timeout=TIMEOUT) as r:
        body = r.read()
    return (json.loads(body) if body else {}), 'Zoho CRM (read-only search)'


def _key(s):
    return re.sub(r'[^A-Za-z0-9]', '', str(s or '')).upper()


def _norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip().lower().strip('.,')


def _digits(s):
    d = re.sub(r'\D', '', str(s or ''))
    return d[-10:] if len(d) >= 10 else d


def _load_record(job):
    p = os.path.join(job, '_facts.json')
    if not os.path.isfile(p):
        return None, 'no _facts.json in the job folder'
    try:
        with open(p, encoding='utf-8') as f:
            data = json.load(f)
    except Exception as e:
        return None, f'_facts.json unreadable ({type(e).__name__})'
    if not isinstance(data, dict):
        return None, '_facts.json is not a JSON object'
    return data, None


def _workup_markup_and_sub(path):
    """(markup_fraction | None, subcontractor_label_found, subcontractor_name)
    from a work-up's visible sheets: the "Markup %" column of the row whose
    first cell reads "Total", and the text beside "Subcontractor Name:"."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    markup, sub_found, sub_name = None, False, ''
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible':
            continue
        hdr = None
        for row in ws.iter_rows():
            cells = [c for c in row if c.value is not None]
            if not cells:
                continue
            for c in cells:
                if isinstance(c.value, str) and c.value.strip().lower() == 'markup %':
                    hdr = hdr or c.column
                if (isinstance(c.value, str) and not sub_found
                        and c.value.strip().lower().startswith('subcontractor name')):
                    sub_found = True
                    inline = c.value.split(':', 1)[1].strip() if ':' in c.value else ''
                    beside = [ws.cell(row=c.row, column=c.column + k).value
                              for k in (1, 2)]
                    beside = [str(v).strip() for v in beside
                              if isinstance(v, str) and v.strip()]
                    sub_name = inline or (beside[0] if beside else '')
            first = cells[0].value
            if (markup is None and hdr and isinstance(first, str)
                    and first.strip().lower() == 'total'):
                v = ws.cell(row=cells[0].row, column=hdr).value
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    markup = float(v)
    return markup, sub_found, sub_name


def _families_table():
    import build_package
    return build_package.system_families()


def parse_system_types(value, table):
    """Zoho System_Types is a multi-select whose options can themselves be
    comma-joined combinations ("Fire Alarm, Intrusion Alarm"). Split every
    option into names and map each through the SOW-BASELINES §3 table.
    Returns (families, unrecognised_names). Display only; never decides."""
    items = value if isinstance(value, list) else (
        [value] if isinstance(value, str) and value else [])
    fams, unknown = [], []
    for item in items:
        for name in re.split(r'[,;]', str(item)):
            name = name.strip()
            if not name:
                continue
            fam = table.get(name)
            if fam is None:
                unknown.append(name)
            elif fam not in fams:
                fams.append(fam)
    return fams, unknown


def run(job, rec, prospect, subscriber, site, monthly_total, workups):
    """Emit the five Zoho findings through ``rec(status, item, detail)``."""
    miss = missing_credentials()
    if miss:
        why = (f'{", ".join(miss)} unset; the Zoho cross-checks need a '
               'read-only Zoho credential (docs/verifier-coverage.md)')
        for c in CHECKS:
            rec('SKIP', c, why)
        return
    if not prospect:
        for c in CHECKS:
            rec('SKIP', c, 'no Prospect # on the schedule to look up')
        return
    try:
        resp, src = _search(prospect)
    except Exception as e:
        for c in CHECKS:
            rec('SKIP', c, f'Zoho unreachable or unreadable ({type(e).__name__})')
        return
    data = resp.get('data') if isinstance(resp, dict) else None
    # The search criterion is not trusted either: keep only exact matches.
    deals = [d for d in (data or []) if isinstance(d, dict)
             and _key(d.get('True_Lead_Number')) == _key(prospect)]
    if len(deals) != 1:
        found = ('no deal' if not deals else
                 f'{len(deals)} deals')
        rec('WARN', C_PROSPECT,
            f'schedule Prospect # "{prospect}" vs Zoho: {found} with that '
            f'True_Lead_Number ({src})')
        for c in CHECKS[1:]:
            rec('SKIP', c, 'no single Zoho deal for the Prospect #')
        return
    deal = deals[0]
    rec('PASS', C_PROSPECT,
        f'schedule Prospect # "{prospect}" vs Zoho True_Lead_Number '
        f'"{deal.get("True_Lead_Number")}" ({src})')

    record, rwhy = _load_record(job)
    _check_contact(rec, deal, record, rwhy)
    _check_deal(rec, deal, subscriber, site, workups)
    _check_master(rec, deal, record, rwhy)
    _check_term(rec, deal, monthly_total)


def _check_contact(rec, deal, record, rwhy):
    if record is None:
        rec('SKIP', C_CONTACT, rwhy)
        return
    cust = record.get('customer') or {}
    p_phones = [str(cust.get(k)) for k in ('phone', 'cell') if cust.get(k)]
    z_phones = [str(deal.get(k)) for k in
                ('Deal_Contact_Office_Phone', 'Deal_Contact_Cell_Phone') if deal.get(k)]
    p_mail = str(cust.get('email') or '').strip()
    z_mail = str(deal.get('Deal_Contact_Email') or '').strip()
    parts, bad, compared = [], False, False
    if p_phones and z_phones:
        compared = True
        ok = bool({_digits(x) for x in p_phones} & {_digits(x) for x in z_phones})
        bad |= not ok
    parts.append(f'phone: record "{" / ".join(p_phones) or "(none)"}" vs Zoho '
                 f'"{" / ".join(z_phones) or "(none)"}"')
    if p_mail and z_mail:
        compared = True
        bad |= p_mail.lower() != z_mail.lower()
    parts.append(f'email: record "{p_mail or "(none)"}" vs Zoho "{z_mail or "(none)"}"')
    detail = '; '.join(parts)
    if not compared:
        rec('SKIP', C_CONTACT, f'nothing to compare, one side empty; {detail}')
    else:
        rec('WARN' if bad else 'PASS', C_CONTACT, detail)


def _check_deal(rec, deal, subscriber, site, workups):
    parts, bad, compared = [], False, False
    z_name = str(deal.get('Deal_Name') or '').strip()
    names = [n for n in (subscriber, site) if n and n.strip()]
    if z_name and names:
        compared = True
        ok = any(_norm(n) in _norm(z_name) for n in names)
        bad |= not ok
        parts.append(f'deal name: package "{" / ".join(names)}" vs Zoho "{z_name}"'
                     + ('' if ok else ' (Zoho name names neither)'))
    else:
        parts.append(f'deal name not compared: package "{" / ".join(names) or "(none)"}"'
                     f' vs Zoho "{z_name or "(none)"}"')

    wu = workups[0] if workups else None
    markup = sub_found = None
    sub_name = ''
    if wu:
        try:
            markup, sub_found, sub_name = _workup_markup_and_sub(wu)
        except Exception as e:
            parts.append(f'work-up unreadable ({type(e).__name__})')
    wu_tag = os.path.basename(wu) if wu else 'no work-up in folder'

    z_markup = deal.get('Project_Markup')
    try:
        z_markup_f = float(z_markup) if z_markup not in (None, '') else None
    except (TypeError, ValueError):
        z_markup_f = None
    if markup is not None and z_markup_f is not None:
        compared = True
        ok = abs(markup * 100 - z_markup_f) <= MARKUP_TOLERANCE_PTS
        bad |= not ok
        parts.append(f'markup: work-up {markup * 100:.2f}% vs Zoho {z_markup_f:.2f}%')
    else:
        parts.append('markup not compared: work-up '
                     f'{"(not found)" if markup is None else f"{markup * 100:.2f}%"}'
                     f' ({wu_tag}) vs Zoho {z_markup!r}')

    z_sub = deal.get('Subcontractor_Used')
    if sub_found and isinstance(z_sub, bool):
        compared = True
        p_sub = bool(sub_name)
        bad |= p_sub != z_sub
        parts.append(f'subcontractor: work-up {"yes" if p_sub else "no"}'
                     f'{f" ({sub_name})" if sub_name else ""} vs Zoho '
                     f'Subcontractor_Used {z_sub}')
    else:
        parts.append('subcontractor not compared: work-up '
                     f'{"label found" if sub_found else "no Subcontractor Name label"}'
                     f' ({wu_tag}) vs Zoho {z_sub!r}')
    detail = '; '.join(parts)
    if not compared:
        rec('SKIP', C_DEAL, detail)
    else:
        rec('WARN' if bad else 'PASS', C_DEAL, detail)


def _check_master(rec, deal, record, rwhy):
    if record is None:
        rec('SKIP', C_MASTER, rwhy)
        return
    import build_package
    try:
        plan = build_package.compose(record)
    except SystemExit as e:
        rec('SKIP', C_MASTER, 'record not composable: '
            + str(e).splitlines()[0][:160])
        return
    family = plan['family']
    box = MASTER_BOXES.get(family)
    if box is None:
        rec('SKIP', C_MASTER, f'no Zoho master-sent checkbox maps to the '
            f'{family} family; nothing to compare')
        return
    situation = (record.get('deal') or {}).get('package_situation')
    expected = bool(plan['agreements'])
    boxes = {b: deal.get(b) for b in MASTER_BOXES.values()}
    bad = bool(boxes[box]) != expected or any(
        bool(v) for b, v in boxes.items() if b != box)
    try:
        fams, unknown = parse_system_types(deal.get('System_Types'),
                                           _families_table())
    except SystemExit:
        fams, unknown = [], []
    zt = (f'; Zoho System_Types {deal.get("System_Types")!r} parses to '
          f'{", ".join(fams) or "(no family)"}'
          + (f', unrecognised {", ".join(unknown)}' if unknown else '')
          + ' (shown, not used to decide)')
    rec('WARN' if bad else 'PASS', C_MASTER,
        f'package: situation "{situation}", {family}, master '
        f'{"in" if expected else "not in"} the package vs Zoho: '
        + ', '.join(f'{b}={v}' for b, v in boxes.items()) + zt)


def _check_term(rec, deal, monthly_total):
    term = deal.get('Effective_Agreement_Term')
    if monthly_total is None:
        rec('SKIP', C_TERM, f'no Monthly Total on the schedule to decide RMR; '
            f'Zoho Effective_Agreement_Term {term!r}')
        return
    try:
        t = float(term) if term not in (None, '') else 0.0
    except (TypeError, ValueError):
        t = 0.0
    detail = (f'schedule Monthly Total ${monthly_total:,.2f} vs Zoho '
              f'Effective_Agreement_Term {term!r} months')
    if monthly_total > 0 and t == 0:
        rec('WARN', C_TERM, detail + '; zero term on a project carrying RMR')
    else:
        rec('PASS', C_TERM, detail)
