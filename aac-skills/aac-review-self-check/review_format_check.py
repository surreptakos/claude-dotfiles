#!/usr/bin/env python3
"""AAC performance review format check. Run this on your draft before you send it.

  python3 review_format_check.py REVIEW.docx --end 7/23/2026 [--start 7/24/2025] [--direct FIRSTNAME] [--no-render]

  --end    required: the last day of the review period
  --start  optional: the first day; checks the span
  --direct optional: the direct's first name, so his own name does not count as an example
  --no-render  skip the page count when LibreOffice is not installed

Requires python-docx (pip install python-docx). The page count needs LibreOffice (soffice) and pdfinfo.
"""
import re, sys, argparse, subprocess, tempfile, os, datetime, shutil
from docx import Document

RATINGS = ["exceeded expectations", "met expectations", "not met expectations", "did not meet expectations"]
RESULTS = ["promotion", "vertical growth", "horizontal growth", "no change", "current role and responsibilities"]
PRESCRIPTIVE = [r"\bneeds? to\b", r"\bshould\b", r"\bmust\b",
                r"\bwould benefit from\b", r"\bought to\b", r"\bis expected to\b", r"\bshall\b"]
ABBR = ["Mr", "Ms", "Mrs", "Dr", "Inc", "Corp", "Co", "Mfg", "St", "Ave", "Blvd", "Jr", "Sr", "vs", "etc", "No", "Dept",
        "a.m", "p.m", "U.S", "e.g", "i.e", "Q1", "Q2", "Q3", "Q4"]
MONTHS = "January|February|March|April|May|June|July|August|September|October|November|December|Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec"
DATE_NUM = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{2,4})\b")
DATE_TXT = re.compile(r"\b(%s)\.? (\d{1,2})(?:st|nd|rd|th)?,? (\d{4})\b" % MONTHS)
STOP = set("""I He She His Her Him It They Them Their The A An This That These Those In On At For To From With By As Of And But Or
Both Two Three Four Five Six Seven Eight Nine Ten Net New Last Next Q1 Q2 Q3 Q4 RMR GP GPM CCTV AAC CRM Guidance Point Strength
Weakness Core Message""".split())


def sentences(text):
    t = text.strip()
    for a in ABBR:
        t = re.sub(r"\b%s\." % re.escape(a), a.replace(".", "<dot>") + "<dot>", t)
    t = re.sub(r"(\d)\.(\d)", r"\1<dot>\2", t)            # decimals
    t = re.sub(r"\$(\d[\d,]*)\.(\d\d)", r"$\1<dot>\2", t)  # cents
    parts = re.split(r"(?<=[.!?])[\"”’)]?\s+(?=[\"“(]?[A-Z0-9$])", t)
    return [p.replace("<dot>", ".") for p in parts if p.strip()]


def parse_date(s):
    m = DATE_NUM.fullmatch(s.strip())
    if m:
        mo, d, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100: y += 2000
        return datetime.date(y, mo, d)
    m = DATE_TXT.fullmatch(s.strip())
    if m:
        mo = [x for x in MONTHS.split("|")].index(m.group(1)) % 12 + 1
        return datetime.date(int(m.group(3)), mo, int(m.group(2)))
    raise ValueError("bad date " + s)


def all_dates(text):
    out = []
    for m in DATE_NUM.finditer(text):
        try: out.append((parse_date(m.group(0)), m.group(0)))
        except ValueError: pass
    for m in DATE_TXT.finditer(text):
        try: out.append((parse_date(m.group(0)), m.group(0)))
        except ValueError: pass
    return out


def has_anchor(text, direct):
    if DATE_NUM.search(text) or DATE_TXT.search(text): return True
    if re.search(r"\$\d", text) or re.search(r"\d+(\.\d+)?%", text): return True
    words = re.findall(r"\b[A-Z][A-Za-z&]+\b", text)
    for i, w in enumerate(words):
        if w in STOP or w == direct or w.rstrip("’'s") == direct: continue
        # sentence-initial words are ambiguous; require the word not to follow a period
        idx = text.find(w)
        before = text[:idx].rstrip()
        if before and before[-1] in ".!?": continue
        if idx == 0: continue
        return True
    return False


def load(path):
    d = Document(path)
    paras = [p.text.strip() for p in d.paragraphs]
    header = {}
    for p in paras:
        m = re.search(r"Dates covered by this review:\s*(.+?)\s*[-–—]\s*(\S+)", p)
        if m: header["start"], header["end"] = m.group(1).strip(), m.group(2).strip()
        m = re.search(r"Delivered by:\s*(.+?)\s+Date:\s*(.+)$", p)
        if m: header["delivered_by"], header["date"] = m.group(1).strip(), m.group(2).strip()
        m = re.search(r"Last Review Date:\s*(\S+)", p)
        if m: header["last_review"] = m.group(1)
    core = next((p[len("Core Message:"):].strip() for p in paras if p.startswith("Core Message:")), "")
    # guidance: after the "Guidance for the next year" line until a signature line
    gi = next((i for i, p in enumerate(paras) if p.lower().startswith("guidance for the next year")), None)
    guidance = []
    if gi is not None:
        for p in paras[gi + 1:]:
            if not p: continue
            if p.startswith("Direct Signature") or p.startswith("Reviewer Signature"): break
            guidance.append(p)
    strengths, weaknesses = [], []
    for t in d.tables:
        hdr = [c.text.strip().lower() for c in t.rows[0].cells]
        if not any("strength" in h for h in hdr): continue
        for r in t.rows[1:]:
            cells = r.cells
            for ci, bucket in ((0, strengths), (1, weaknesses)):
                if ci >= len(cells): continue
                for p in cells[ci].paragraphs:
                    if p.text.strip(): bucket.append(p.text.strip())
    return header, core, strengths, weaknesses, guidance, "\n".join(paras + strengths + weaknesses)


def page_count(path):
    if not shutil.which("soffice"): return None
    out = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out, path],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
    pdf = [f for f in os.listdir(out) if f.endswith(".pdf")]
    if not pdf: return None
    info = subprocess.run(["pdfinfo", os.path.join(out, pdf[0])], capture_output=True, text=True).stdout
    m = re.search(r"Pages:\s+(\d+)", info)
    return int(m.group(1)) if m else None


def format_check(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("docx"); ap.add_argument("--end", required=True); ap.add_argument("--start")
    ap.add_argument("--direct", default=""); ap.add_argument("--no-render", action="store_true")
    a = ap.parse_args(argv)
    end = parse_date(a.end)
    header, core, S, W, G, alltext = load(a.docx)
    fixes, passes = [], []

    # 1 header period
    hs, he = header.get("start"), header.get("end")
    try:
        he_d = parse_date(he) if he else None
    except ValueError:
        he_d = None
    if he_d != end:
        fixes.append(f"Header: dates covered end {he or '[missing]'}; must end {end.strftime('%-m/%-d/%Y')}.")
    else:
        passes.append("Header end date")
    if a.start:
        st = parse_date(a.start)
        try:
            hs_d = parse_date(hs) if hs else None
        except ValueError:
            hs_d = None
        if hs_d != st:
            fixes.append(f"Header: dates covered start {hs or '[missing]'}; must start {st.strftime('%-m/%-d/%Y')}.")
        else:
            passes.append("Header start date")
    notes = []
    if not header.get("date") or "[" in header.get("date", "") or "<" in header.get("date", ""):
        notes.append("Header: Date field is a placeholder; it carries the delivery date when the review goes out.")

    # 2 core message
    cs = sentences(core)
    lc = core.lower()
    if len(cs) > 3:
        fixes.append(f"Core Message: {len(cs)} sentences; must be three or fewer.")
    r_hits = [r for r in RATINGS if r in lc]
    if "not met expectations" in lc or "did not meet expectations" in lc:
        r_hits = [r for r in r_hits if r != "met expectations"]
    if len(r_hits) != 1:
        fixes.append("Core Message: Rating phrase missing or doubled; must contain exactly one of exceeded expectations, met expectations, not met expectations.")
    res_hits = [r for r in RESULTS if r in lc]
    if len(res_hits) == 0:
        fixes.append("Core Message: Result phrase missing; must name promotion, vertical growth, horizontal growth, or maintaining current role and responsibilities.")
    elif len(set(res_hits) - {"no change", "current role and responsibilities"}) > 1:
        fixes.append("Core Message: more than one Result named; must name one.")
    if not fixes or all(not f.startswith("Core Message") for f in fixes):
        passes.append("Core Message form")

    # 3 strengths and weaknesses
    for label, items in (("S", S), ("W", W)):
        for i, txt in enumerate(items, 1):
            name = f"{label}{i}"
            n = len(sentences(txt))
            if n not in (2, 4):
                fixes.append(f"{name}: {n} sentences; must be two (Sum-Ex) or four (SEER).")
            if not has_anchor(txt, a.direct):
                fixes.append(f"{name}: no date, figure, or named account or person; must contain one.")
            for pat in PRESCRIPTIVE:
                m = re.search(pat, txt, re.I)
                if m:
                    fixes.append(f"{name}: contains \"{m.group(0)}\"; not allowed in a Strength or Weakness (write the behavior; an instruction belongs in Guidance).")
                    break
    if not any(re.match(r"[SW]\d+:", f) for f in fixes):
        passes.append(f"Strengths and Weaknesses form ({len(S)} S, {len(W)} W)")

    # 4 second person anywhere
    yous = re.findall(r"\b[Yy]ou(r|rs|rself)?\b", alltext)
    if yous:
        fixes.append(f"Voice: \"you\" or \"your\" appears {len(yous)} time(s); the review is third person about the direct.")
    else:
        passes.append("Third person")
    # the reviewer is first person; "his manager" and the like refer to the reviewer in third person.
    # edge case: a direct who manages people may legitimately have someone else called "his manager."
    mgr = re.findall(r"\b(?:his|her|their|the) (?:manager|supervisor)\b", alltext, re.I)
    if mgr:
        fixes.append(f"Voice: \"{mgr[0]}\" refers to the reviewer in third person; the review is first person for the reviewer.")
    else:
        passes.append("First person for the reviewer")

    # 5 dates after the period end
    late = sorted({s for dt, s in all_dates(alltext) if dt > end and s != header.get("date")})
    if late:
        fixes.append(f"Dates: {', '.join(late)} fall after the period end {end.strftime('%-m/%-d/%Y')}; nothing in the review is dated after it.")
    else:
        passes.append("No dates after period end")

    # 6 guidance
    if not G:
        fixes.append("Guidance: no Guidance points found under \"Guidance for the next year.\"")
    for i, g in enumerate(G, 1):
        m = re.match(r"Guidance Point (\d+):\s*(\S+)", g)
        if not m:
            fixes.append(f"Guidance {i}: does not start \"Guidance Point N:\" followed by an action verb.")
            continue
        first = m.group(2).rstrip(",.").lower()
        if first in ("you", "your", "he", "she", "the", "a", "an", "it", "this") or (a.direct and first == a.direct.lower()):
            fixes.append(f"Guidance Point {m.group(1)}: starts with \"{m.group(2)}\"; must start with an action verb.")
        n = len(sentences(g))
        if n > 3:
            fixes.append(f"Guidance Point {m.group(1)}: {n} sentences; must be one to three.")
    if not any(f.startswith("Guidance") for f in fixes):
        passes.append(f"Guidance form ({len(G)} points)")

    # 7 one page
    pc = None if a.no_render else page_count(a.docx)
    if pc is None:
        print("(page count not checked)", file=sys.stderr)
    elif pc > 1:
        fixes.append(f"Length: {pc} pages; the review is one page.")
    else:
        passes.append("One page")

    print("FORMAT CHECK:", "PASS" if not fixes else f"FAIL ({len(fixes)} fix{'es' if len(fixes) != 1 else ''})")
    for f in fixes: print("  -", f)
    print("Passed:", "; ".join(passes))
    for n in notes: print("Note:", n)
    return 0 if not fixes else 1


def main():
    return format_check(sys.argv[1:])


if __name__ == "__main__":
    sys.exit(main())
